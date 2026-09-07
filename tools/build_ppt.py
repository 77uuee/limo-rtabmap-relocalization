# -*- coding: utf-8 -*-
"""slides/*.dc.html 4장을 편집 가능한 PowerPoint(.pptx)로 재구성."""
import re, os
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SRC = "/Users/paddone/limo-lab/3d map/slides"
OUT = "/Users/paddone/limo-lab/3d map/slides/rtabmap-reliability.pptx"

# 원본 캔버스가 1280x720 px 이므로 96dpi 기준 1px = 9525 EMU 로 그대로 옮긴다.
# (13.333in x 7.5in = 표준 16:9 슬라이드와 정확히 일치)
PX = 9525
def E(px): return Emu(int(round(px * PX)))
def S(px): return Pt(px * 0.75)          # CSS px -> pt

FONT = "Apple SD Gothic Neo"

C = dict(blue=RGBColor(0x2f,0x5f,0xe0), red=RGBColor(0xd9,0x2c,0x2c),
         green=RGBColor(0x2e,0x9e,0x5b), amber=RGBColor(0xe2,0xa6,0x3d),
         amber_d=RGBColor(0xb0,0x7d,0x1a),
         ink=RGBColor(0x1a,0x1a,0x1a), body=RGBColor(0x33,0x33,0x33),
         mute=RGBColor(0x66,0x66,0x66), mute2=RGBColor(0x88,0x88,0x88),
         mute3=RGBColor(0x55,0x55,0x55), mute4=RGBColor(0x77,0x77,0x77),
         white=RGBColor(0xff,0xff,0xff), line=RGBColor(0xdd,0xdd,0xdd),
         tline=RGBColor(0xd8,0xde,0xe9), gray=RGBColor(0xcc,0xcc,0xcc))
BG = dict(blue=RGBColor(0xf4,0xf7,0xff), green=RGBColor(0xf2,0xfb,0xf5),
          red=RGBColor(0xfd,0xf3,0xf3), amber=RGBColor(0xfd,0xf9,0xf0),
          gray=RGBColor(0xf8,0xf8,0xf8), th=RGBColor(0xee,0xf2,0xfc))

prs = Presentation()
prs.slide_width, prs.slide_height = E(1280), E(720)
BLANK = prs.slide_layouts[6]


# ---------- 인라인 마크업(<b>, <br>) 파서 ----------
def emit(tf, markup, size, color, *, bold=False, line_spacing=1.35,
         align=PP_ALIGN.LEFT, space_after=0):
    """'a<b>b</b>c' + '<br>' 을 문단/런으로 풀어 텍스트프레임에 채운다."""
    markup = markup.replace("&nbsp;", " ")
    first_para = True
    for lineno, line in enumerate(markup.split("<br>")):
        p = tf.paragraphs[0] if first_para else tf.add_paragraph()
        first_para = False
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        for seg in re.split(r"(<b>.*?</b>)", line):
            if not seg:
                continue
            m = re.fullmatch(r"<b>(.*?)</b>", seg, re.S)
            r = p.add_run()
            r.text = m.group(1) if m else seg
            f = r.font
            f.name, f.size, f.color.rgb = FONT, S(size), color
            f.bold = True if m else bold
    return tf


def text(sl, x, y, w, h, markup, size, color, **kw):
    tb = sl.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = kw.pop("anchor", MSO_ANCHOR.TOP)
    emit(tf, markup, size, color, **kw)
    return tb


def box(sl, x, y, w, h, fill, border, radius=14):
    sh = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, E(x), E(y), E(w), E(h))
    sh.adjustments[0] = radius / min(w, h)      # 반지름을 px 단위로 맞춤
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if border is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = border; sh.line.width = Pt(1.1)
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh


def badge(sl, x, y, label, fill):
    w = 20 * 2 + len(label) * 18            # 좌우 padding 20px + 글자폭 근사
    sh = box(sl, x, y, w, 38, fill, None, radius=10)
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    emit(tf, label, 17, C['white'], bold=True, align=PP_ALIGN.CENTER)
    return x + w


def header(sl, y, label, badge_fill, title_markup, title_size=27):
    xe = badge(sl, 52, y, label, badge_fill)
    text(sl, xe + 16, y - 2, 1228 - xe, 46, title_markup, title_size, C['ink'],
         bold=True, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0)


def pic(sl, x, y, w, h, name):
    p = sl.shapes.add_picture(os.path.join(SRC, name), E(x), E(y), E(w), E(h))
    p.line.color.rgb = C['line']; p.line.width = Pt(0.75)
    return p


def caption(sl, x, y, w, markup, size=14, ls=1.35):
    text(sl, x, y, w, 44, markup, size, C['mute'],
         align=PP_ALIGN.CENTER, line_spacing=ls)


def boxtitle(sh, markup, size, color, pad_l=22, pad_t=16):
    """도형 위에 제목+본문을 얹기 위한 내부 텍스트프레임 여백 설정."""
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left, tf.margin_right = E(pad_l), E(pad_l)
    tf.margin_top, tf.margin_bottom = E(pad_t), E(pad_t)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    return tf


def para(tf, markup, size, color, *, bold=False, ls=1.35, before=0, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.line_spacing = ls
    p.space_before = Pt(before)
    for seg in re.split(r"(<b>.*?</b>)", markup.replace("&nbsp;", " ")):
        if not seg:
            continue
        m = re.fullmatch(r"<b>(.*?)</b>", seg, re.S)
        r = p.add_run(); r.text = m.group(1) if m else seg
        f = r.font; f.name, f.size, f.color.rgb = FONT, S(size), color
        f.bold = True if m else bold


# ============================ 1 · 과제와 방법 ============================
s = prs.slides.add_slide(BLANK)
header(s, 44, "과제", C['blue'],
       "만든 지도를 <b>언제까지 믿을 수 있나</b>", 33)
# 제목 뒷부분만 빨강으로: 런 색상 개별 지정
tb = s.shapes[-1]
tb.text_frame.paragraphs[0].runs[1].font.color.rgb = C['red']

text(s, 52, 96, 1176, 24,
     "카메라 전용 RTAB-Map 지도에서 재위치추정 성공률 실측 · LIMO Pro + Orbbec RGB-D · 2026.08.30 ~ 09.02",
     16, C['mute'])

LX, LW, RX, RW = 52, 547, 627, 601

b = box(s, LX, 136, LW, 366, BG['blue'], C['blue'])
tf = boxtitle(b, "", 0, None, 22, 18)
para(tf, "측정 방법", 18, C['blue'], bold=True, first=True)
for line in [
    "<b>지도 1회 고정</b> — 매핑 후 원본 DB 읽기 전용 보관, 측정은 매회 사본으로 (cold start)",
    "<b>정답 = 바닥 테이프</b> — 지점마다 위치·방향 테이프 등록, 시스템 출력과 독립된 기준",
    "<b>판정 허용치</b> — 위치 0.5 m · 방향 30° 이내면 성공, 벗어나면 오탐, 미발행이면 미수렴",
    "<b>2단계 프로토콜</b> — 정지 20초 → 미수렴 시 저속 회전, 총 60초",
    "<b>반복</b> — 지점 × 조건당 3회, 스크립트 자동 기록 (수렴 시각·오차·믿음 점프)",
]:
    para(tf, line, 15.5, C['body'], ls=1.45, before=11)

b = box(s, LX, 516, LW, 160, BG['green'], C['green'])
tf = boxtitle(b, "", 0, None, 22, 16)
para(tf, "조건 (조도 중심)", 18, C['green'], bold=True, first=True)
para(tf, "<b>C1</b> 기준 — 주간 · 조명 켬 (매핑과 동일)", 15.5, C['body'], ls=1.5, before=9)
para(tf, "<b>C3</b> 소등 — 주간 · 조명 전체 소등 (창측 자연광 잔존)", 15.5, C['body'], ls=1.5)
para(tf, "※ 시간대·배치 조건은 지도교수 지침에 따라 조도 조건으로 통합", 14, C['mute4'], ls=1.5)

pic(s, RX, 210, RW, 343, "points_map.jpg")
caption(s, RX, 564, RW,
        "측정 지점 7곳 테이프 등록 (강의실 P1·P2·P3·P7, 복도 P4·P5·P6) — 이 중 <b>P1·P2·P4·P5·P7</b> 5곳 실측")


# ============================== 2 · 결과 ==============================
s = prs.slides.add_slide(BLANK)
header(s, 40, "결과", C['blue'],
       "소등해도 성공률은 같다 — <b>C1 42% ≈ C3 40%</b>")
s.shapes[-1].text_frame.paragraphs[0].runs[1].font.color.rgb = C['blue']

LX, LW, RX, RW = 52, 660, 740, 488

rows = [
    ("지점", "C1 기준 (불 켬)", "C3 소등 (주간)", "특성", None),
    ("<b>P1</b> 강의실 중앙", "<b>3/3</b> (0.5s · 3cm)", "<b>3/3</b> (0.6s · 8cm)", "기준값 — 안정", ('g', 'g')),
    ("<b>P2</b> 창문 벽", "0/3", "1/3 (즉시 · 2cm)", "특징점 빈곤", ('r', None)),
    ("<b>P4</b> 복도 초입", "1/3 (즉시 · 13cm)", "1/3 (즉시 · 7cm)", "확률적", (None, None)),
    ("<b>P5</b> 복도 중간", "미측정*", "1/3 (즉시 · 6cm)", "확률적", (None, None)),
    ("<b>P7</b> 북쪽 경계", "1/3 (즉시 · 9cm)", "0/3", "오인 다발", (None, 'r')),
    ("<b>전체</b>", "<b>5/12 (42%)</b>", "<b>6/15 (40%)</b>", "—", (None, None)),
]
TW = [190, 175, 168, 127]
RH = 46
gt = s.shapes.add_table(len(rows), 4, E(LX), E(96), E(LW), E(RH * len(rows))).table
for i, w in enumerate(TW):
    gt.columns[i].width = E(w)
gt.first_row = False          # 기본 밴딩/헤더 스타일 끄기 (직접 칠한다)
gt.horz_banding = False

for ri, row in enumerate(rows):
    gt.rows[ri].height = E(RH)
    tint = row[4]
    for ci, cell_md in enumerate(row[:4]):
        c = gt.cell(ri, ci)
        c.margin_left = c.margin_right = E(8)
        c.margin_top = c.margin_bottom = E(4)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        c.fill.solid()
        if ri == 0:
            c.fill.fore_color.rgb = BG['th']
        elif tint and ci in (1, 2) and tint[ci - 1]:
            c.fill.fore_color.rgb = BG['green'] if tint[ci - 1] == 'g' else BG['red']
        else:
            c.fill.fore_color.rgb = C['white']
        tf = c.text_frame
        tf.word_wrap = True
        emit(tf, cell_md, 15, C['blue'] if ri == 0 else C['body'],
             bold=(ri == 0), align=PP_ALIGN.CENTER, line_spacing=1.1)

text(s, LX, 428, LW, 22,
     "성공 = 60초 내 테이프 정답 대비 0.5 m·30° 이내 수렴 · *P5 C1은 카메라 고장·배터리로 미측정",
     13, C['mute2'])

b = box(s, LX, 460, LW, 128, BG['blue'], C['blue'], radius=12)
tf = boxtitle(b, "", 0, None, 18, 14)
para(tf, "성패를 가르는 것은 조명이 아니라 ① <b>지점의 특징점 품질</b> ② <b>cold start 확률성</b> "
         "(성공은 대부분 기동 직후 정지 상태에서, 회전 중엔 모션 블러로 실패) ③ <b>시각적 혼동</b>",
     15, C['body'], ls=1.5, first=True)

IW, IH = 353, 265
IX = RX + (RW - IW) // 2
pic(s, IX, 96, IW, IH, "p2fail.jpg")
caption(s, RX, 368, RW, "P2 카메라 시점 — 흰 벽·흰 가전뿐, 잡을 특징점이 없다", 13.5)
pic(s, IX, 398, IW, IH, "lightsoff.jpg")
caption(s, RX, 670, RW, "C3 소등 상태의 카메라 시점 — 자연광 + 자동 노출로 장면 유지", 13.5)


# ========================= 3 · 실패 지점 지도 =========================
s = prs.slides.add_slide(BLANK)
header(s, 40, "실패 지점 지도", C['blue'],
       "실패는 무작위가 아니라 <b>특정 지점에 몰린다</b>")
s.shapes[-1].text_frame.paragraphs[0].runs[1].font.color.rgb = C['red']

LX, LW, RX, RW = 52, 638, 718, 510
pic(s, LX, 100, LW, 364, "fail_map.jpg")
tb = text(s, LX, 476, LW, 44,
          "지점별 성공률 (C1+C3 합산) — 초록 안정 · 빨강 취약 · 회색 미측정",
          14, C['mute'], align=PP_ALIGN.CENTER)
for r in tb.text_frame.paragraphs[0].runs:      # 단일 런이므로 색 강조는 분리 생성
    pass
tb.text_frame.clear()
emit(tb.text_frame, "지점별 성공률 (C1+C3 합산) — ", 14, C['mute'], align=PP_ALIGN.CENTER)
p = tb.text_frame.paragraphs[0]
for txt, col, bold in [("초록", C['green'], True), (" 안정 · ", C['mute'], False),
                       ("빨강", C['red'], True), (" 취약 · 회색 미측정", C['mute'], False)]:
    r = p.add_run(); r.text = txt
    r.font.name, r.font.size, r.font.color.rgb, r.font.bold = FONT, S(14), col, bold

b = box(s, RX, 100, RW, 158, BG['red'], C['red'])
tf = boxtitle(b, "", 0, None, 20, 16)
para(tf, "시각적 혼동 (perceptual aliasing)", 17, C['red'], bold=True, first=True)
para(tf, "P1↔P7 오인이 <b>양방향으로 4회</b> 재현됨. 책상 반복 무늬가 두 지점에서 비슷하게 보여, "
         "로봇이 P7에 서서 “나는 P1”이라 믿는 식의 <b>4 m급 믿음 점프</b>가 기록됨 (전부 자동 검출).",
     15, C['body'], ls=1.5, before=9)

b = box(s, RX, 270, RW, 158, BG['amber'], C['amber'])
tf = boxtitle(b, "", 0, None, 20, 16)
para(tf, "취약 지점의 공통점", 17, C['amber_d'], bold=True, first=True)
para(tf, "<b>P2</b> 흰 벽·가전 — 특징점 빈곤 (회전 스윕에도 미수렴)", 15, C['body'], ls=1.5, before=9)
para(tf, "<b>P7</b> 개활 구역 — 반복 무늬 + 관측 부족 방향", 15, C['body'], ls=1.5)
para(tf, "<b>P4·P5</b> 복도 — 첫 회차만 성공하는 확률성", 15, C['body'], ls=1.5)

b = box(s, RX, 440, RW, 140, BG['gray'], C['gray'])
tf = boxtitle(b, "", 0, None, 20, 16)
para(tf, "오탐(엉뚱한 곳으로 수렴) 방지를 위해 매칭 문턱을 올린 상태 (MinInliers 20→30). "
         "오인은 줄었지만 어려운 지점의 성공률이 함께 낮아짐 — <b>정확도와 가용성의 트레이드오프</b>를 그대로 보고.",
     15, C['mute3'], ls=1.5, first=True)


# ============================== 4 · 결론 ==============================
s = prs.slides.add_slide(BLANK)
header(s, 40, "결론", C['red'],
       "지도의 유효기간은 시간이 아니라 <b>“배치가 유지되는 동안”</b>")
s.shapes[-1].text_frame.paragraphs[0].runs[1].font.color.rgb = C['red']

LX, LW, RX, RW = 52, 601, 681, 547

b = box(s, LX, 100, LW, 282, BG['blue'], C['blue'])
tf = boxtitle(b, "", 0, None, 20, 16)
para(tf, "실증 사례 — 의자 사건 (9/2 오전)", 17, C['blue'], bold=True, first=True)
para(tf, "<b>9/1</b>&nbsp;&nbsp;기준 지점 P1, 3회 연속 즉시 수렴 (오차 3 cm)", 15, C['body'], ls=1.45, before=10)
para(tf, "<b>9/2 오전</b>&nbsp;&nbsp;밤사이 의자들이 <b>수십 cm씩만</b> 이동", 15, C['body'], ls=1.45, before=8)
para(tf, "<b>→</b>&nbsp;&nbsp;같은 지점·같은 조명에서 <b>2회 연속 미수렴</b>.", 15, C['body'], ls=1.45, before=8)
para(tf, "    후보 매칭은 되지만 실물이 달라 기하검증 전멸", 15, C['body'], ls=1.45)
para(tf, "<b>→</b>&nbsp;&nbsp;<b>5분 이어매핑</b>으로 지도 갱신 후 0.4초 만에 수렴 복구", 15, C['body'], ls=1.45, before=8)

b = box(s, LX, 394, LW, 252, BG['green'], C['green'])
tf = boxtitle(b, "", 0, None, 20, 16)
para(tf, "운영 시사점", 17, C['green'], bold=True, first=True)
para(tf, "① 조도 변화(주간 소등)에는 강하다 — 성공률 동일", 15, C['body'], ls=1.55, before=9)
para(tf, "② 가구 배치 변화에는 약하다 — “살짝”으로도 기준 지점 상실", 15, C['body'], ls=1.55)
para(tf, "③ 복구는 저렴하다 — 재매핑이 아니라 <b>5분 이어매핑</b>이면 충분", 15, C['body'], ls=1.55)
para(tf, "④ 기하(형상) 기반 LiDAR+AMCL에는 없는 유형의 실패로 추정", 15, C['body'], ls=1.55)
para(tf, "    → 동일 지점 직접 비교는 향후 과제", 15, C['body'], ls=1.55)

pic(s, RX, 100, RW, 410, "chair.jpg")
caption(s, RX, 522, RW,
        "미수렴 당시 로봇 시점(P1) — 근거리 의자가 시야를 지배.<br>사람 눈에는 “거의 그대로”인 배치가 카메라에는 다른 장소",
        14, ls=1.4)

prs.save(OUT)
print("saved:", OUT, os.path.getsize(OUT), "bytes,", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
